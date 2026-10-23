"""Ingestão da documentação de regras de negócio: Volume -> chunks -> Delta -> Vector Search.

Estrutura no Volume: /Volumes/<catalogo>/<schema>/documentos/<id_do_agente>/*.pdf|docx|md|txt
Só agentes com `documentos: true` são indexados.
"""
from __future__ import annotations

import hashlib
import os
import time
from typing import Callable, Dict, List

from .config import AppConfig

EXTENSOES = ("pdf", "docx", "md", "txt", "csv", "json", "html")


def pasta_documentos(cfg: AppConfig) -> str:
    return f"/Volumes/{cfg.projeto.catalogo}/{cfg.projeto.schema_}/documentos"


def preparar_volumes(spark, cfg: AppConfig) -> None:
    cat, sch = cfg.projeto.catalogo, cfg.projeto.schema_
    spark.sql(f"CREATE SCHEMA IF NOT EXISTS {cat}.{sch}")
    for vol in ("documentos", "llmops"):
        spark.sql(f"CREATE VOLUME IF NOT EXISTS {cat}.{sch}.{vol}")


def contar_documentos(pasta: str) -> int:
    if not os.path.isdir(pasta):
        return 0
    return sum(
        1 for _, _, arqs in os.walk(pasta) for a in arqs if a.lower().rsplit(".", 1)[-1] in EXTENSOES
    )


def extrair_texto(caminho: str) -> str:
    ext = caminho.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader

        return "\n".join((p.extract_text() or "") for p in PdfReader(caminho).pages)
    if ext == "docx":
        import docx

        d = docx.Document(caminho)
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:  # regras costumam vir em tabelas
            for linha in t.rows:
                partes.append(" | ".join(c.text.strip() for c in linha.cells))
        return "\n".join(partes)
    if ext in EXTENSOES:
        with open(caminho, encoding="utf-8", errors="ignore") as f:
            return f.read()
    return ""


def ler_documentos(cfg: AppConfig, volume: str, log: Callable[[str], None] = print) -> List[Dict[str, str]]:
    ids = {e.id for e in cfg.especialistas if e.documentos}
    docs = []
    for eid in sorted(ids):
        pasta = os.path.join(volume, eid)
        if not os.path.isdir(pasta):
            log(f"⚠️  Sem pasta de documentos para '{eid}': {pasta}")
            continue
        for raiz, _, arquivos in os.walk(pasta):
            for nome in arquivos:
                caminho = os.path.join(raiz, nome)
                texto = extrair_texto(caminho).strip()
                if texto:
                    docs.append({"especialista": eid, "fonte": nome, "texto": texto})
                else:
                    log(f"⚠️  Sem texto extraído (formato não suportado ou PDF escaneado): {caminho}")
    return docs


def gerar_chunks(docs: List[Dict[str, str]], chunk_size: int = 1000, chunk_overlap: int = 150) -> List[Dict[str, str]]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=chunk_overlap, separators=["\n\n", "\n", ". ", " ", ""]
    )
    linhas = []
    for doc in docs:
        for i, trecho in enumerate(splitter.split_text(doc["texto"])):
            chave = f"{doc['especialista']}|{doc['fonte']}|{i}"
            linhas.append({
                "chunk_id": hashlib.md5(chave.encode(), usedforsecurity=False).hexdigest(),  # só identificador
                "especialista": doc["especialista"],
                "fonte": doc["fonte"],
                "conteudo": trecho,
            })
    return linhas


def ingerir_documentos(
    spark,
    cfg: AppConfig,
    volume: str | None = None,
    chunk_size: int = 1000,
    chunk_overlap: int = 150,
    log: Callable[[str], None] = print,
) -> Dict[str, int]:
    if not cfg.usa_documentos:
        log("Nenhum agente usa documentação — ingestão ignorada.")
        return {}
    vs_cfg = cfg.vector_search
    volume = volume or pasta_documentos(cfg)
    preparar_volumes(spark, cfg)

    docs = ler_documentos(cfg, volume, log)
    linhas = gerar_chunks(docs, chunk_size, chunk_overlap)
    log(f"📄 {len(docs)} documento(s) → {len(linhas)} trecho(s)")
    if not linhas:
        log(f"⚠️  Nenhum documento encontrado em {volume}/<id_do_agente>/. Suba os arquivos e rode de novo.")
        return {}

    spark.sql(f"""
        CREATE TABLE IF NOT EXISTS {vs_cfg.tabela_origem} (
            chunk_id STRING NOT NULL, especialista STRING, fonte STRING, conteudo STRING
        ) TBLPROPERTIES (delta.enableChangeDataFeed = true)
    """)
    df = spark.createDataFrame(linhas, "chunk_id string, especialista string, fonte string, conteudo string")
    df.write.mode("overwrite").option("overwriteSchema", "false").saveAsTable(vs_cfg.tabela_origem)

    from databricks.vector_search.client import VectorSearchClient

    vsc = VectorSearchClient(disable_notice=True)
    if vs_cfg.endpoint not in [e["name"] for e in vsc.list_endpoints().get("endpoints", [])]:
        log(f"Criando endpoint de Vector Search '{vs_cfg.endpoint}' (pode levar ~10 min)…")
        vsc.create_endpoint_and_wait(name=vs_cfg.endpoint, endpoint_type="STANDARD")
    try:
        indice = vsc.get_index(vs_cfg.endpoint, vs_cfg.indice_padrao)
        log("Índice existente — sincronizando…")
        indice.sync()
    except Exception:  # noqa: BLE001
        log(f"Criando índice {vs_cfg.indice_padrao}…")
        indice = vsc.create_delta_sync_index_and_wait(
            endpoint_name=vs_cfg.endpoint,
            index_name=vs_cfg.indice_padrao,
            source_table_name=vs_cfg.tabela_origem,
            primary_key="chunk_id",
            pipeline_type="TRIGGERED",
            embedding_source_column="conteudo",
            embedding_model_endpoint_name=vs_cfg.modelo_embedding,
            columns_to_sync=vs_cfg.colunas,
        )
    while not indice.describe().get("status", {}).get("ready", False):
        log("Aguardando o índice ficar pronto…")
        time.sleep(20)

    por_agente: Dict[str, int] = {}
    for l in linhas:
        por_agente[l["especialista"]] = por_agente.get(l["especialista"], 0) + 1
    log(f"✅ Índice pronto: {por_agente}")
    return por_agente
